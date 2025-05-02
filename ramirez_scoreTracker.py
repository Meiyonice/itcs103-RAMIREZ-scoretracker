from openpyxl import Workbook, load_workbook
import tkinter as tk
from tkinter import messagebox
from openpyxl.styles import Font
import os

# ========== Submit Data Function ==========
def submit_data():
    name = name_entry.get().title()
    score = score_entry.get()

    if not name or not score:
        messagebox.showerror("Input Error", "All fields are required!")
        return

    try:
        score = int(score)
    except ValueError:
        messagebox.showerror("Input Error", "Score must be a valid numerical value.")
        return
    
    if not os.path.exists("student_scores.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Scores"
        ws.append(["Name", "Score", "Remarks"])

        for cell in ws[1]:
            cell.font = Font(bold=True)

        wb.save("student_scores.xlsx")
    else:
        wb = load_workbook("student_scores.xlsx")
        ws = wb["Scores"]
    
    for row in reversed(range(2, ws.max_row + 1)):
        if ws.cell(row=row, column=1).value == "Average Score:":
            ws.delete_rows(row - 1, 3)
            break

    remarks = "Pass" if score >= 75 else "Fail"

    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == name:
            ws.cell(row=row, column=2, value=score)
            ws.cell(row=row, column=3, value=remarks)
            break
    else:
        ws.append([name, score, remarks])

    scores = [row[1].value for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=False) if isinstance(row[1].value, (int, float))]

    if scores:
        avg_score = sum(scores) / len(scores)
        avg_remarks = "Pass" if avg_score >= 75 else "Fail"

        ws.append([" ", " ", " "])
        ws.append(["Average Score:", avg_score, avg_remarks])

    wb.save("student_scores.xlsx")

    messagebox.showinfo("Success", "Data saved successfully!")

    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)

# ========== View Data Function ==========
def view_data():
    try:
        wb = load_workbook("student_scores.xlsx")
        ws = wb["Scores"]
    except FileNotFoundError:
        messagebox.showerror("File Not Found", "The file 'student_scores.xlsx' does not exist.")
        return
    except KeyError:
        messagebox.showerror("Sheet Not Found", "The sheet 'Scores' is missing in the Excel file.")
        return
    
    view_window = tk.Toplevel(window)
    view_window.title("Scores")

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        for j, value in enumerate(row):
            if i == 0:
                label = tk.Label(view_window, text=str(value), padx=5, pady=2, font=("Helvetica", 10, "bold"))
            else:
                label = tk.Label(view_window, text=str(value), padx=5, pady=2)
            
            label.grid(row=i, column=j, sticky="nsew", padx=1, pady=1)

# ========== Tkinter UI ==========
window = tk.Tk()
window.title("Ramirez Score Tracker")
window.geometry("300x200")
window.resizable(False, False)

tk.Label(window, text="Name:", font=("Helvetica", 10, "bold")).grid(row=0, column=0, padx=10, pady=15)
tk.Label(window, text="Score:", font=("Helvetica", 10, "bold")).grid(row=1, column=0, padx=10)

name_entry = tk.Entry(window, width=30)
score_entry = tk.Entry(window, width=30)

name_entry.grid(row=0, column=1)
score_entry.grid(row=1, column=1)

tk.Button(window, text="Submit Data", command=submit_data, width=10, bg="#FF7F50", fg="white").grid(row=3, column=1, pady=25)
tk.Button(window, text="View Data", command=view_data, width=10, bg="#008080", fg="white").grid(row=4, column=1)

window.mainloop()